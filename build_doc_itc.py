"""
Generatore DOCX - Documento di Informazione alle Ditte
Fascicolo 2511/25 Ord. 2 - Ditta I.T.C. S.r.l.
Revisione n.2 verricelli Bacino BRIN
Template replicato da Fascicolo 6301/26 Mancarella
"""
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

ASSETS = "/projects/sandbox/ITALIC/_assets"
OUT = "/projects/sandbox/ITALIC/Doc_info_ditte_Fascicolo_2511_25_ITC_Verricelli_BRIN.docx"

doc = Document()

# --- PAGE SETUP ---
for section in doc.sections:
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(1.35)
    section.bottom_margin = Cm(1.75)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)

# --- STYLES ---
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(12)


# --- HELPER FUNCTIONS ---
def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top','left','bottom','right'):
        if edge in kwargs:
            element = OxmlElement(f'w:{edge}')
            element.set(qn('w:val'), kwargs[edge].get('val','single'))
            element.set(qn('w:sz'), kwargs[edge].get('sz','4'))
            element.set(qn('w:color'), kwargs[edge].get('color','000000'))
            tcBorders.append(element)
    tcPr.append(tcBorders)

def shade_cell(cell, color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), color)
    shd.set(qn('w:val'), 'clear')
    tcPr.append(shd)

def add_heading1(text):
    p = doc.add_paragraph(text, style='Heading 1')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for run in p.runs:
        run.font.name = 'Calibri'
    return p

def add_heading2(text):
    p = doc.add_paragraph(text, style='Heading 2')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.font.name = 'Calibri'
    return p

def add_body(text, bold=False, italic=False, size=12):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    return p

def add_body_center(text, bold=False, italic=False, size=12):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    return p


# --- HEADER (tabella 3 colonne come template) ---
header = doc.sections[0].header
header.is_linked_to_previous = False
# rimuovi paragrafo vuoto default
for p in header.paragraphs:
    p.clear()

htbl = header.add_table(rows=1, cols=3, width=Cm(18))
htbl.autofit = True
c0 = htbl.cell(0, 0)
c1 = htbl.cell(0, 1)
c2 = htbl.cell(0, 2)

p0 = c0.paragraphs[0]
r0 = p0.add_run("Redatto a cura del:\n  Servizio Prevenzione e Protezione\nArsenale M. M. Taranto")
r0.font.size = Pt(8)
r0.font.name = 'Calibri'

p1 = c1.paragraphs[0]
p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
r1 = p1.add_run("DOCUMENTO DI \nINFORMAZIONE ALLE DITTE\n- D. Lgs. 81/2008 Art. 26 comma 1 lett. b) -")
r1.font.size = Pt(9)
r1.font.name = 'Calibri'
r1.bold = True

p2 = c2.paragraphs[0]
p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r2 = p2.add_run("_______ 2025")
r2.font.size = Pt(9)
r2.font.name = 'Calibri'

# --- FOOTER ---
footer = doc.sections[0].footer
footer.is_linked_to_previous = False
for p in footer.paragraphs:
    p.clear()
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = fp.add_run("SERVIZIO PREVENZIONE E PROTEZIONE \u2013 ARSENALE M.M. TARANTO")
fr.font.size = Pt(8)
fr.font.name = 'Calibri'


# --- FRONTESPIZIO ---
# Logo (image1.png dal template - stemma MM)
logo_path = os.path.join(ASSETS, "image1.png")
if os.path.exists(logo_path):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    r.add_picture(logo_path, width=Cm(2.5))

doc.add_paragraph()  # spazio

add_body_center("ARSENALE MILITARE MARITTIMO", bold=True, italic=True, size=16)
add_body_center("TARANTO", bold=True, italic=True, size=16)

# spazi
for _ in range(3):
    doc.add_paragraph()

add_body_center("DOCUMENTO DI INFORMAZIONE", bold=True, size=19)
add_body_center("ALLE DITTE APPALTATRICI", bold=True, size=19)
add_body_center("ai sensi dell'art. 26, comma 1, lett. b)", size=12)
add_body_center("D. Lgs. 81/2008 e s.m.i.", size=12)

doc.add_paragraph()
doc.add_paragraph()

add_body_center("Oggetto", bold=True, size=12)
add_body("PROGRAMMA DI MANTENIMENTO CONDIZIONI OPERATIVE (MCO) PER I BACINI DI "
         "CARENAGGIO DI MARINARSEN TARANTO \u2013 REVISIONE N. 2 VERRICELLI DEL BACINO BRIN "
         "\u2013 DITTA I.T.C. S.r.l.", bold=True, size=12)

doc.add_paragraph()
doc.add_paragraph()

add_body("Il Responsabile del Servizio Prevenzione e Protezione", bold=True, italic=True)
add_body("FST ing. Giancarlo CAFORIO", bold=True, italic=True)

doc.add_page_break()


# --- PREMESSA ---
add_heading1("PREMESSA")

add_body(
    "Lo scopo del presente documento \u00e8 quello di fornire le informazioni, ai sensi "
    "dell\u2019art. 26 comma 1 lettera b) del D. Lgs. 81/2008 e s.m.i., relative ai luoghi "
    "di lavoro oggetto del contratto e dipendenti dall\u2019organizzazione del Committente, "
    "al fine di consentire alla Ditta appaltatrice di adottare le misure di prevenzione "
    "e protezione pi\u00f9 idonee in relazione ai rischi specifici presenti nell\u2019ambiente "
    "in cui il personale \u00e8 destinato ad operare."
)

add_body(
    "Ai sensi dell\u2019art. 26 comma 3 D.Lgs. 81/2008, la Ditta Appaltatrice \u00e8 tenuta "
    "a partecipare alle riunioni di coordinamento indette dal Committente e a prendere "
    "visione del DUVRI (Appendice A alla S.T.) ove predisposto. Il presente documento "
    "costituisce parte integrante della documentazione contrattuale (Fascicolo 2511/25 "
    "\u2013 Ordine nr. 2) e dovr\u00e0 essere portato a conoscenza, da parte del Datore di Lavoro "
    "della Ditta appaltatrice, di tutto il personale impiegato nelle lavorazioni, ivi "
    "inclusi eventuali subappaltatori autorizzati e lavoratori autonomi.",
    italic=True
)

doc.add_paragraph()


# --- NUOVO PARAGRAFO POST-PREMESSA ---
add_heading1("INFORMAZIONI SUI RISCHI INTERFERENZIALI E SUL COORDINAMENTO OPERATIVO")

add_body(
    "Le attivit\u00e0 oggetto del presente affidamento si svolgeranno all\u2019interno del "
    "comprensorio dell\u2019Arsenale Militare Marittimo di Taranto, in area Bacino Brin, "
    "ove possono essere presenti, in concomitanza o in prossimit\u00e0, ulteriori attivit\u00e0 "
    "lavorative gestite secondo diverse linee di comando e distinti regimi "
    "tecnico-organizzativi."
)

add_body(
    "Il presente documento fornisce pertanto le informazioni di carattere generale sui "
    "rischi presenti nell\u2019ambiente di lavoro e sui rischi interferenziali conoscibili e "
    "stabili al momento della sua redazione, restando esclusa la descrizione dettagliata "
    "delle lavorazioni variabili in corso nel cantiere navale militare e nel cantiere "
    "pozzo prosciugamento RTL, che saranno oggetto di specifica informazione e "
    "coordinamento da parte del Responsabile Tecnico dei Lavori / Responsabile Esecuzione."
)

add_body(
    "In particolare, il personale della Ditta ITC S.r.l. dovr\u00e0 operare esclusivamente "
    "nelle aree autorizzate dal DEC e secondo le modalit\u00e0 concordate in sede di "
    "coordinamento, evitando ogni interferenza con le attivit\u00e0 militari e con le "
    "lavorazioni di altre imprese eventualmente presenti nel comprensorio."
)

add_body(
    "Per i rischi derivanti da attivit\u00e0 soggette a D.Lgs. 272/1999, il coordinamento "
    "operativo, la convocazione delle riunioni e la formalizzazione per iscritto delle "
    "misure di prevenzione e protezione sono demandati al RTL / Responsabile Esecuzione "
    "dei lavori, fermo restando l\u2019obbligo del committente di aggiornare il presente "
    "documento in caso di nuove interferenze non previste."
)

add_body(
    "Qualsiasi variazione logistica, organizzativa o temporale che possa incidere sulla "
    "sicurezza delle attivit\u00e0 dovr\u00e0 essere comunicata tempestivamente al DEC e, "
    "comunque, prima della ripresa delle lavorazioni interessate."
)

add_body(
    "Si d\u00e0 atto che il presente documento informativo \u00e8 integrato dal documento gi\u00e0 "
    "trasmesso al Responsabile Tecnico dei Lavori / Responsabile Esecuzione (RTL/R.E.) "
    "per le attivit\u00e0 soggette al D.Lgs. 272/1999, allegato al presente fascicolo quale "
    "parte integrante e sostanziale, ai fini del coordinamento delle interferenze e "
    "della formalizzazione delle misure di prevenzione e protezione."
)

doc.add_paragraph()


# --- RIFERIMENTI AL CONTRATTO ---
add_heading1("RIFERIMENTI AL CONTRATTO")

add_body(
    "Fascicolo: 2511/25 \u2013 Programma di mantenimento delle condizioni operative (MCO) "
    "per i bacini di carenaggio di MARINARSEN Taranto \u2013 Revisione n. 2 verricelli del "
    "Bacino BRIN."
)
add_body("Ordine: nr. 2")
add_body("Ditta appaltatrice: I.T.C. S.r.l.")
add_body("Committente: Arsenale Militare Marittimo di Taranto \u2013 MARINARSEN TARANTO")

doc.add_paragraph()

# --- AREE E LOCALI ---
add_heading1("AREE E LOCALI DOVE POSSONO ESSERE SVOLTI I LAVORI")

add_body(
    "Le attivit\u00e0 di revisione dei n. 2 verricelli si svolgeranno nelle seguenti aree "
    "del comprensorio di MARINARSEN Taranto:"
)
add_body("\u2022 Bacino di carenaggio BRIN \u2013 piano di banchina, ciglio bacino, postazioni "
         "dei verricelli di manovra;")
add_body("\u2022 Aree limitrofe ai verricelli necessarie per la movimentazione di componenti, "
         "attrezzature e ricambi;")
add_body("\u2022 Eventuali locali tecnici annessi al bacino, indicati di volta in volta dal DEC/RTL;")
add_body("\u2022 Percorsi di accesso e transito all\u2019interno del comprensorio, dal varco di "
         "ingresso autorizzato fino al Bacino BRIN.")

add_body(
    "Sono espressamente esclusi dall\u2019area di lavoro autorizzata: gli specchi acquei e "
    "le banchine non strettamente connesse alle lavorazioni; i bacini diversi dal BRIN; "
    "le aree in cui sono presenti unit\u00e0 navali militari; le aree militari operative, i "
    "depositi, i locali con regime di accesso riservato; ogni altra zona non "
    "espressamente autorizzata dal DEC."
)

add_body(
    "L\u2019accesso alle aree di lavoro deve avvenire esclusivamente attraverso i percorsi "
    "indicati dal DEC e nei tempi concordati con il RTL/R.E."
)

doc.add_paragraph()


# --- FIGURE RESPONSABILI ---
add_heading1("Figure Responsabili ai sensi del D. Lgs. 81/2008")

add_body(
    "Marinarsen ha una struttura gerarchica definita con Decreto Ministeriale (DM) e "
    "disposizioni interne. Di seguito vengono riportati i nominativi del Datore di "
    "lavoro, dei dirigenti e preposti interessati alla presente attivit\u00e0:"
)

doc.add_paragraph()
add_body("Datore di Lavoro - Direttore protempore dell\u2019Arsenale di Taranto:", bold=True)
add_body("Amm. Isp. Alessandro BATTAGLIA")

doc.add_paragraph()
add_body("Figure di particolare interesse per l\u2019oggetto dei lavori:")
add_body("Capo Reparto Supporto Tecnico Arsenale: Dirigente Tommaso COVIELLO tel.099 775 2356")
add_body("Direttore Operativo (D.O.): F.S.T. Ing. Antonio SURANO tel. 099 775 2603")
add_body("Responsabile Servizio Prevenzione e Protezione: FST Giancarlo CAFORIO tel.099 775 3103")
add_body("Medico competente: Dott. Giovanni TRIA tel.099 775 4028")
add_body("Capo Sezione Sanitaria: C.C. Dott. Valerio SCARANO CATANZARO tel. 099 775 2366")

doc.add_paragraph()
add_body("RAPPRESENTANTI DEI LAVORATORI")
add_body("Personale Civile:")
add_body("Alessandro SCIALPI")
add_body("Giovanni CHIFFI")
add_body("Pierluigi PAULI")
add_body("Pietro AVELLINO")
add_body("Saverio CAPODIFERRO")
add_body("Valentina FALCONE")
add_body("Personale Militare:")
add_body("--")

doc.add_page_break()


# --- 1) INFORMAZIONI GENERALI ---
add_heading1("1) INFORMAZIONI GENERALI SULL\u2019ARSENALE DI TARANTO")

add_body("Denominazione Ente: ARSENALE MILITARE MARITTIMO TARANTO (detto anche MARINARSEN TARANTO)")
doc.add_paragraph()
add_body("Ubicazione:   TARANTO,  Piazza Ammiraglio Pasquale  Leonardi  CAP:  74123")
doc.add_paragraph()
add_body("Organo di Vigilanza Competente:   MARIVIGILANZA \u2013 AREA SUD \u2013 TARANTO")

doc.add_page_break()

# --- 2) EMERGENZE ---
add_heading1("2) INFORMAZIONI RELATIVE ALLA GESTIONE DELLE EMERGENZE, EVACUAZIONE e PRIMO SOCCORSO.")

add_heading2("EVACUAZIONE")
add_body(
    "Di seguito la planimetria con centri di raccolta di tutto il comprensorio "
    "Arsenalizio."
)

# Inserisco planimetria se disponibile (image2.png dal template)
plan_path = os.path.join(ASSETS, "image2.png")
if os.path.exists(plan_path):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    r.add_picture(plan_path, width=Cm(15))

doc.add_paragraph()


# --- CHIAMATE DI EMERGENZA ---
add_body_center("CHIAMATE DI EMERGENZA", bold=True, size=18)
doc.add_paragraph()
add_body_center("EMERGENZA INCENDIO: 115 Vigili del Fuoco", bold=True)
doc.add_paragraph()
add_body_center("EMERGENZA SANITARIA: 118 Pronto Soccorso")
doc.add_paragraph()
add_body_center("SALA MEDICA ARSENALE tel. 099 775 2841", bold=True)
add_body_center("Primo Soccorso per eventi non gravi solo durante i seguenti orari:", bold=True)
doc.add_paragraph()
add_body_center("luned\u00ec \u2013 gioved\u00ec dalle 08:00 alle 16:30", bold=True)
add_body_center("venerd\u00ec dalle 08:00 alle 12:00", bold=True)

doc.add_paragraph()

# --- EMERGENZA ELETTRICA (box) ---
tbl_ee = doc.add_table(rows=1, cols=1)
tbl_ee.style = 'Table Grid'
cell_ee = tbl_ee.cell(0, 0)
p_ee = cell_ee.paragraphs[0]
p_ee.alignment = WD_ALIGN_PARAGRAPH.LEFT
r_ee = p_ee.add_run(
    "EMERGENZA ELETTRICA\n\n"
    "Se necessita un intervento in emergenza sulla rete elettrica:\n"
    "Guardia Circuiti - Sezione Reti Elettriche:\n"
    "da telefono militare: 2601\n"
    "da telefono civile: 099 775 2601\n\n"
    "Elettricista di turno H24 (per bacini e banchine):\n"
    "da telefono militare: 2601\n"
    "da telefono civile: 099 775 2601"
)
r_ee.font.name = 'Calibri'
r_ee.font.size = Pt(11)
r_ee.bold = True

doc.add_paragraph()


# --- EMERGENZA BACINO BRIN ---
add_body_center("EMERGENZA BACINO BRIN", bold=True, size=15)
doc.add_paragraph()

add_body(
    "In considerazione della specificit\u00e0 del Bacino BRIN, la Ditta deve attenersi "
    "alle seguenti disposizioni di emergenza:"
)
add_body(
    "\u2022 Caduta di persona o di materiali in bacino: interrompere le attivit\u00e0, dare "
    "l\u2019allarme alla Centrale Operativa e al DEC, non tentare il recupero in autonomia "
    "se non in condizioni di sicurezza, agevolare l\u2019intervento delle squadre di soccorso."
)
add_body(
    "\u2022 Allagamento accidentale o variazione anomala del livello d\u2019acqua: allontanarsi "
    "immediatamente dal piano vasca e dalle gallerie/pozzetti, raggiungere il piano di "
    "banchina lungo i percorsi pi\u00f9 brevi e segnalati, dare l\u2019allarme."
)
add_body(
    "\u2022 Emergenza meccanica del verricello / dei sistemi di tenuta della porta-bacino / "
    "dei sistemi di pompaggio: arrestare le manovre in sicurezza, segnalare l\u2019anomalia "
    "al DEC e al RTL/R.E., non riavviare le manovre fino a verifica e autorizzazione "
    "formale."
)
add_body(
    "\u2022 Sversamenti di oli, idrocarburi o sostanze pericolose: contenere lo sversamento "
    "con i mezzi a disposizione (kit antinquinamento di cantiere) e dare immediato "
    "avviso, evitando lo scarico in mare."
)
add_body(
    "\u2022 Incendio in bacino o nei locali tecnici annessi: utilizzare i mezzi di "
    "estinzione disponibili solo se addestrati e in sicurezza, attivare l\u2019allarme, "
    "evacuare verso il punto di raccolta indicato."
)
add_body(
    "L\u2019attivazione delle pompe del bacino, la manovra della porta-bacino e ogni "
    "operazione di prosciugamento o riempimento sono di competenza esclusiva del "
    "personale di MARINARSEN. La Ditta non pu\u00f2 in alcun caso intervenire su tali sistemi."
)

doc.add_page_break()


# --- ACCESSI ---
add_body_center("INFORMAZIONI SUGLI ACCESSI NEL COMPRENSORIO", bold=True, size=15)
doc.add_paragraph()

# Box ordine del giorno accessi
tbl_acc = doc.add_table(rows=1, cols=1)
tbl_acc.style = 'Table Grid'
cell_acc = tbl_acc.cell(0, 0)
p_acc = cell_acc.paragraphs[0]
r_acc = p_acc.add_run(
    "Ordine del giorno n\u00b0___ del __/__/2025\n"
    "Si applicano le vigenti disposizioni di MARINARSEN Taranto in materia di accesso "
    "al comprensorio per il personale delle Ditte appaltatrici."
)
r_acc.font.name = 'Calibri'
r_acc.font.size = Pt(11)

doc.add_paragraph()

add_body_center("REPERIBILIT\u00c0 TELEFONICA SERVIZI DI GUARDIA E DI VIGILANZA", bold=True, size=15)
add_body("Comunicazione di Servizio VDA/___ del __/__/2025", bold=True)

doc.add_paragraph()

add_body(
    "L\u2019accesso al comprensorio dell\u2019Arsenale \u00e8 regolato da specifiche disposizioni "
    "interne. La Ditta \u00e8 tenuta a trasmettere preventivamente al DEC l\u2019elenco "
    "nominativo del personale, richiedere il rilascio dei titoli di accesso (badge), "
    "transitare esclusivamente dai varchi indicati e restituire i titoli al termine "
    "delle attivit\u00e0."
)

add_body(
    "All\u2019interno del comprensorio \u00e8 fatto obbligo di: rispettare la segnaletica; "
    "procedere con i mezzi a velocit\u00e0 ridotta (max 20 km/h), dando precedenza ai mezzi "
    "militari, ai mezzi di emergenza e ai pedoni; non eseguire riprese fotografiche "
    "senza autorizzazione; non introdurre materiali non autorizzati."
)

doc.add_page_break()


# --- 3) INFORMAZIONI SUI RISCHI ---
add_heading1("3) INFORMAZIONI  SUI RISCHI")

add_heading2("CODIFICAZIONE FATTORI DI RISCHIO")

# Tabella codifica (come nel template)
risk_codes = [
    ("4 \u2013 Movimentazione manuale dei carichi",
     "Movimentazione di componenti meccanici dei verricelli (tamburi, riduttori, "
     "freni, perni). Rischio muscolo-scheletrico durante smontaggio e rimontaggio."),
    ("7 \u2013 Rumore / 8 \u2013 Vibrazioni",
     "Utilizzo di utensili elettrici/pneumatici (smerigliatrici, avvitatori ad "
     "impulso, martelli) durante la revisione. Rumore e vibrazioni mano-braccio."),
    ("9 \u2013 App. a pressione",
     "Eventuale utilizzo di aria compressa per pulizia componenti, prove di tenuta "
     "circuiti idraulici dei verricelli."),
    ("10 \u2013 Apparecchi di sollevamento",
     "Utilizzo di gru, paranchi, argani per sollevamento e movimentazione organi "
     "meccanici dei verricelli (tamburi, motoriduttori, freni). Carichi sospesi."),
    ("11 \u2013 App. trasp. mov. interna",
     "Transito di autogrù, carrelli elevatori, mezzi di trasporto componenti nelle "
     "aree di banchina e bacino. Interferenza con personale a piedi."),
    ("12 \u2013 Attrezzature manuali",
     "Utilizzo di chiavi dinamometriche, estrattori, martinetti, attrezzi manuali "
     "durante le operazioni di smontaggio/rimontaggio dei verricelli."),
    ("13 \u2013 Impianti elettrici",
     "Interventi su quadri di alimentazione e motori elettrici dei verricelli. "
     "Rischio elettrocuzione. Lavori elettrici ai sensi CEI 11-27."),
    ("14 \u2013 Macchine",
     "Organi meccanici in movimento dei verricelli (tamburi, ingranaggi, freni). "
     "Rischio di cesoiamento, trascinamento, schiacciamento durante prove funzionali."),
    ("15 \u2013 Sorgenti d\u2019incendio e/o esplosione",
     "Lavorazioni a caldo (saldatura, molatura) su componenti metallici dei "
     "verricelli. Presenza di oli e grassi lubrificanti."),
    ("17 \u2013 Locali/attrezzature di interconnessione",
     "Strade, banchine, ciglio bacino, scale di accesso. Interferenze con traffico "
     "militare e civile. Banchine non delimitate: pericolo caduta in mare."),
    ("21 \u2013 Condizioni climatiche",
     "Attivit\u00e0 all\u2019aperto su banchina e ciglio bacino: esposizione a intemperie, "
     "vento forte, superfici rese scivolose dalla pioggia."),
    ("23 \u2013 Organizzazione del lavoro",
     "Coordinamento con DEC e RTL/R.E. per ogni fase. Gestione delle interferenze "
     "con altre imprese e attivit\u00e0 militari. Delimitazione area di lavoro."),
]

tbl_r = doc.add_table(rows=1 + len(risk_codes), cols=2)
tbl_r.style = 'Table Grid'
# intestazione
hdr0 = tbl_r.cell(0, 0)
hdr1 = tbl_r.cell(0, 1)
hdr0.paragraphs[0].add_run("Codice FR (DVR Rev.14/2024)").bold = True
hdr1.paragraphs[0].add_run("Ambiti di potenziale rischio").bold = True
shade_cell(hdr0, 'D9E2F3')
shade_cell(hdr1, 'D9E2F3')

for i, (code, desc) in enumerate(risk_codes, start=1):
    tbl_r.cell(i, 0).paragraphs[0].add_run(code).font.size = Pt(10)
    tbl_r.cell(i, 1).paragraphs[0].add_run(desc).font.size = Pt(10)

doc.add_paragraph()


# --- QUADRO DI SINTESI ---
add_heading2("QUADRO DI SINTESI")

# 4 - MMC
add_body("4 \u2013 Movimentazione Manuale dei Carichi", bold=True, italic=True, size=13)
doc.add_paragraph()
add_body(
    "Le operazioni di smontaggio e rimontaggio dei verricelli comportano la "
    "movimentazione di componenti meccanici di peso significativo (tamburi, "
    "motoriduttori, freni, perni, boccole)."
)
add_body(
    "Per componenti di peso superiore a 25 kg: utilizzare obbligatoriamente mezzi "
    "meccanici di sollevamento (paranchi, gru, martinetti). Adottare posture corrette "
    "e alternare gli operatori per limitare l\u2019esposizione."
)

doc.add_paragraph()

# 7/8 - Rumore/Vibrazioni
add_body("7 \u2013 Rumore / 8 \u2013 Vibrazioni", bold=True, italic=True, size=13)
doc.add_paragraph()
add_body(
    "L\u2019utilizzo di smerigliatrici, avvitatori ad impulso, martelli e utensili "
    "pneumatici durante la revisione genera livelli di rumore e vibrazioni significativi."
)
add_body(
    "Dotare il personale di protezioni uditive (cuffie o tappi EN 352, SNR \u2265 25 dB) "
    "durante le lavorazioni rumorose. Valutare l\u2019esposizione a vibrazioni mano-braccio "
    "ai sensi del D.Lgs. 81/08 Titolo VIII Capo III. Alternare gli operatori."
)

doc.add_paragraph()

# 10 - Sollevamento
add_body("10 \u2013 Apparecchi di Sollevamento", bold=True, italic=True, size=13)
doc.add_paragraph()
add_body(
    "Il sollevamento e la movimentazione di organi meccanici dei verricelli (tamburi, "
    "riduttori, freni) mediante gru di banchina, autogrù o paranchi costituiscono "
    "un\u2019attivit\u00e0 a rischio elevato per la presenza di carichi sospesi in prossimit\u00e0 "
    "del ciglio bacino."
)
add_body(
    "\u00c8 vietata la sosta di personale sotto carichi sospesi. Delimitare l\u2019area di "
    "tiro. Verificare l\u2019idroneit\u00e0 delle brache, delle funi e dei ganci prima di ogni "
    "sollevamento. Le attrezzature di sollevamento devono essere sottoposte a verifiche "
    "periodiche in corso di validit\u00e0. Coordinare le manovre con il personale di "
    "MARINARSEN preposto alla gru di banchina."
)

doc.add_paragraph()

# 11 - Trasporto
add_body("11 \u2013 Apparecchi Trasporto / Movimentazione Interna", bold=True, italic=True, size=13)
doc.add_paragraph()
add_body(
    "L\u2019autogrù e i mezzi di trasporto della Ditta transitano nelle aree interne "
    "dell\u2019Arsenale in prossimit\u00e0 di personale e strutture."
)
add_body(
    "Limite di velocit\u00e0 20 km/h nel comprensorio, 10 km/h nelle aree operative. "
    "Delimitare l\u2019area di manovra con barriere e segnaletica. Il personale a terra "
    "deve indossare indumenti ad alta visibilit\u00e0 (EN ISO 20471 Classe 2). "
    "Non sostare con i mezzi nei pressi di uscite di emergenza e vie di fuga."
)


doc.add_paragraph()

# 13 - Impianti elettrici
add_body("13 \u2013 Impianti elettrici", bold=True, italic=True, size=13)
doc.add_paragraph()
add_body(
    "Gli interventi sui motori elettrici e sui quadri di alimentazione dei verricelli "
    "comportano rischio di elettrocuzione. I lavori elettrici devono essere eseguiti da "
    "personale qualificato PES/PAV ai sensi della CEI 11-27."
)
add_body(
    "Prima di ogni intervento: verificare il sezionamento e il blocco dell\u2019alimentazione "
    "(procedura LOTO). Non \u00e8 consentito effettuare lavori sotto tensione senza specifica "
    "autorizzazione scritta. Qualsiasi apparecchiatura elettrica introdotta dovr\u00e0 "
    "rispondere alle normative vigenti e alle norme CEI applicabili."
)
add_body(
    "Per qualsiasi intervento che coinvolga l\u2019impianto elettrico di MARINARSEN \u00e8 "
    "obbligatoria l\u2019assistenza del personale abilitato dell\u2019Arsenale (PES/PAV)."
)

doc.add_paragraph()

# 14 - Macchine
add_body("14 \u2013 Macchine", bold=True, italic=True, size=13)
doc.add_paragraph()
add_body(
    "Durante le prove funzionali dei verricelli revisionati, gli organi meccanici in "
    "movimento (tamburi, ingranaggi, freni a nastro/disco) presentano rischio di "
    "cesoiamento, trascinamento e schiacciamento."
)
add_body(
    "Le prove funzionali devono essere eseguite con area sgombra e delimitata, sotto "
    "la supervisione del RTL/R.E. e del preposto della Ditta. Ripristinare tutte le "
    "protezioni (carter, griglie) prima delle prove in movimento. Nessun operatore "
    "deve trovarsi nel raggio d\u2019azione delle funi e dei tamburi durante il funzionamento."
)

doc.add_paragraph()

# 15 - Incendio
add_body("15 \u2013 Sorgenti d\u2019incendio e/o esplosione", bold=True, italic=True, size=13)
doc.add_paragraph()
add_body(
    "Le lavorazioni a caldo (saldatura, taglio, molatura) su componenti metallici "
    "dei verricelli e la presenza di oli e grassi lubrificanti generano rischio "
    "incendio."
)
add_body(
    "Richiedere obbligatoriamente il \u201cPermesso di Lavoro a Caldo\u201d al DEC prima "
    "dell\u2019inizio di saldature, tagli o molature. Allontanare materiali combustibili "
    "dal raggio di 10 m. Garantire la presenza di un estintore a polvere o CO\u2082 in "
    "prossimit\u00e0 del punto di lavoro. Mantenere la vigilanza antincendio per almeno "
    "60 minuti dopo la cessazione dei lavori a caldo."
)


doc.add_paragraph()

# 17 - Interconnessione
add_body("17 \u2013 Locali/attrezzature di Interconnessione", bold=True, italic=True, size=13)
doc.add_paragraph()
add_body(
    "Il comprensorio presenta strade, banchine, marciapiedi, scale e sottoservizi. "
    "Le banchine sul mare e il ciglio del Bacino BRIN non sono ovunque delimitate: "
    "pericolo di caduta in mare e in bacino."
)
add_body(
    "Rispettare i percorsi stradali disposti. Non ingombrare con la propria "
    "attrezzatura le vie di evacuazione, non impedire l\u2019accesso ai presidi "
    "antincendio e ai quadri elettrici. Prestare attenzione durante il transito "
    "veicolare per la presenza di beole rotte/lesionate. Velocit\u00e0 massima veicolare "
    "consentita: 20 km/h."
)
add_body(
    "Le banchine sul mare e il ciglio del bacino non sono delimitate. Pericolo di "
    "caduta in acqua. In prossimit\u00e0 del ciglio bacino privo di parapetti: utilizzo "
    "obbligatorio di imbracatura anticaduta (EN 361) assicurata a punto fisso."
)

doc.add_paragraph()

# 21 - Condizioni climatiche
add_body("21 \u2013 Condizioni Climatiche", bold=True, italic=True, size=13)
doc.add_paragraph()
add_body(
    "Le attivit\u00e0 si svolgono prevalentemente all\u2019aperto (banchina e ciglio bacino). "
    "Esposizione a intemperie, vento forte, superfici rese scivolose dalla pioggia."
)
add_body(
    "In caso di vento forte (> 60 km/h): sospendere le operazioni di sollevamento "
    "con gru/autogrù. In caso di pioggia intensa: valutare la sospensione delle "
    "attivit\u00e0 in prossimit\u00e0 del ciglio bacino. Durante le operazioni estive, "
    "garantire pause adeguate e idratazione."
)

doc.add_paragraph()

# 23 - Organizzazione del lavoro
add_body("23 \u2013 Organizzazione del lavoro", bold=True, italic=True, size=13)
doc.add_paragraph()
add_body(
    "Il personale della Ditta deve essere identificabile mediante tesserino di "
    "riconoscimento con fotografia, formato ed aggiornato alla sicurezza con i corsi "
    "previsti dalla legge e informato su quanto prescritto nel presente documento."
)
add_body(
    "Delimitare l\u2019area di intervento con barriere fisiche e idonea cartellonistica. "
    "La recinzione deve impedire l\u2019accesso non autorizzato e proteggere i soggetti "
    "in transito."
)
add_body(
    "Il preposto della Ditta deve garantire l\u2019osservanza degli obblighi di legge, "
    "dell\u2019uso dei DPI e della segnalazione delle inosservanze. Segnalare "
    "tempestivamente al DEC ogni anomalia, quasi-incidente o condizione di pericolo."
)
add_body(
    "Qualsiasi variazione logistica, organizzativa o temporale deve essere comunicata "
    "al DEC prima della ripresa delle lavorazioni. Per le attivit\u00e0 ricadenti "
    "nell\u2019ambito del D.Lgs. 272/1999, il coordinamento \u00e8 demandato al RTL/R.E."
)

doc.add_page_break()


# --- TABELLA DPI ---
add_body("Riepilogo DPI per Rischio \u2013 Fascicolo 2511/25 Ditta I.T.C. S.r.l.", bold=True)
doc.add_paragraph()

dpi_data = [
    ("4 \u2013 Movim. Manuale Carichi",
     "Guanti antitaglio; calzature di sicurezza con puntale (EN ISO 20345); "
     "cintura lombare per carichi > 25 kg; ausili meccanici obbligatori."),
    ("7/8 \u2013 Rumore / Vibrazioni",
     "Cuffie o tappi auricolari (EN 352, SNR \u2265 25 dB) durante utilizzo "
     "smerigliatrici, avvitatori ad impulso e utensili pneumatici."),
    ("9 \u2013 App. a pressione",
     "Occhiali di protezione (EN 166); guanti; calzature di sicurezza."),
    ("10 \u2013 App. di sollevamento",
     "Elmetto (EN 397); calzature di sicurezza; indumenti ad alta visibilit\u00e0 "
     "(EN ISO 20471); guanti da lavoro."),
    ("11 \u2013 Trasporto / Movimentazione",
     "Indumenti ad alta visibilit\u00e0 (EN ISO 20471 Classe 2); calzature di "
     "sicurezza; elmetto nelle aree sotto carichi sospesi."),
    ("12 \u2013 Attrezzature manuali",
     "Guanti antitaglio; occhiali di protezione; calzature di sicurezza."),
    ("13 \u2013 Impianti Elettrici",
     "Guanti isolanti (EN 60903); calzature isolanti (EN 50321); occhiali; "
     "indumenti non conduttivi. Solo personale PES/PAV."),
    ("14 \u2013 Macchine",
     "Elmetto; occhiali antiproiezione; guanti antitaglio; calzature di "
     "sicurezza; indumenti aderenti (no parti svolazzanti)."),
    ("15 \u2013 Incendio / Lavori a caldo",
     "Visiera per saldatura (EN 175); guanti per saldatore (EN 12477); "
     "grembiule ignifugo; calzature di sicurezza; estintore a portata di mano."),
    ("17 \u2013 Interconnessione / Caduta in acqua",
     "Imbracatura anticaduta (EN 361) con longe assicurata a punto fisso per "
     "operazioni in prossimit\u00e0 ciglio bacino/banchine non delimitate; calzature "
     "antiscivolo; elmetto."),
    ("21 \u2013 Condizioni climatiche",
     "Indumenti antipioggia/antivento; calzature antiscivolo; crema solare "
     "in periodo estivo; fornitura acqua potabile."),
    ("23 \u2013 Organizzazione",
     "Tesserino identificativo; DPI di base sempre indossati in cantiere "
     "(elmetto, calzature S3, indumenti alta visibilit\u00e0)."),
]

tbl_d = doc.add_table(rows=1 + len(dpi_data), cols=2)
tbl_d.style = 'Table Grid'
h0 = tbl_d.cell(0, 0)
h1 = tbl_d.cell(0, 1)
h0.paragraphs[0].add_run("Rischio").bold = True
h1.paragraphs[0].add_run("DPI obbligatori").bold = True
shade_cell(h0, 'D9E2F3')
shade_cell(h1, 'D9E2F3')

for i, (risk, dpi) in enumerate(dpi_data, start=1):
    tbl_d.cell(i, 0).paragraphs[0].add_run(risk).font.size = Pt(10)
    tbl_d.cell(i, 1).paragraphs[0].add_run(dpi).font.size = Pt(10)

doc.add_page_break()


# --- NOTA FINALE / NATURA INFORMATIVA ---
add_body(
    "NOTA: Il presente documento ha natura informativa e riguarda i rischi conoscibili "
    "e stabili presenti nelle aree di lavoro al momento della sua redazione. Non "
    "sostituisce la valutazione dei rischi propria della Ditta appaltatrice, n\u00e9 il "
    "documento di coordinamento per le attivit\u00e0 ricadenti sotto il D.Lgs. 272/1999. "
    "Le lavorazioni variabili in corso nel cantiere navale militare e nel cantiere "
    "pozzo-prosciugamento RTL sono e restano oggetto di specifica informazione e "
    "coordinamento da parte del RTL/R.E.",
    bold=True, italic=True, size=11
)

doc.add_paragraph()
doc.add_paragraph()

# --- FIRME / PRESA VISIONE ---
add_body(
    "Il sottoscritto, in qualit\u00e0 di Datore di Lavoro / Legale Rappresentante della "
    "Ditta I.T.C. S.r.l., dichiara di aver ricevuto e preso visione del presente "
    "Documento di Informazione alle Ditte, di averne compreso i contenuti e di "
    "accettarne integralmente le prescrizioni. Si impegna a trasmetterne il contenuto "
    "a tutto il proprio personale impiegato.",
    size=11
)

doc.add_paragraph()
doc.add_paragraph()

# Tabella firme
tbl_f = doc.add_table(rows=5, cols=4)
tbl_f.style = 'Table Grid'

# Intestazione
for ci, htext in enumerate(["Soggetto", "Nominativo", "Data", "Firma"]):
    c = tbl_f.cell(0, ci)
    c.paragraphs[0].add_run(htext).bold = True
    shade_cell(c, 'D9E2F3')

rows_data = [
    ("Datore di Lavoro / Legale Rappresentante Ditta I.T.C. S.r.l.", "", "", ""),
    ("RSPP Ditta I.T.C. S.r.l.", "", "", ""),
    ("Referente di cantiere / Preposto Ditta", "", "", ""),
    ("DEC \u2013 MARINARSEN Taranto", "", "", ""),
]
for ri, (s, n, d, f) in enumerate(rows_data, start=1):
    tbl_f.cell(ri, 0).paragraphs[0].add_run(s).font.size = Pt(10)
    tbl_f.cell(ri, 1).paragraphs[0].add_run(n)
    tbl_f.cell(ri, 2).paragraphs[0].add_run(d)
    tbl_f.cell(ri, 3).paragraphs[0].add_run(f)

doc.add_paragraph()
doc.add_paragraph()

# Firma RSPP Committente
add_body("Per il Committente \u2013 MARINARSEN Taranto", bold=True)
doc.add_paragraph()
add_body("Il Responsabile del Servizio di Prevenzione e Protezione")
add_body("FST ing. Giancarlo CAFORIO")
doc.add_paragraph()
add_body("Data: ___/___/______        Firma: ___________________________")

# --- SAVE ---
doc.save(OUT)
print(f"\nDocumento generato con successo:\n  {OUT}")
