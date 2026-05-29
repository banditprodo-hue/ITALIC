# -*- coding: utf-8 -*-
"""
Genera il DOCUMENTO DI INFORMAZIONE ALLE DITTE (art. 26 c.1 lett. b D.Lgs. 81/08)
per le attivita' presso il BACINO FERRATI dell'Arsenale Militare Marittimo di Taranto:
- Fasc. 2511/25 Ord. 1 : taglio e rimozione barche porta GP65 e GP59
- Fasc. 2509/25 Ord. 6 : pulizia platea bacino da fanghi e acque
Struttura mutuata dal Fascicolo 2508.25 (stesso format SPP Marinarsen) e arricchita
secondo lo stile del Documento Informativo Mancarella (Fasc. 6301/26).
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = "/projects/sandbox/ITALIC"
OUT = os.path.join(BASE, "Documento_Informativo_Ditte_Bacino_Ferrati_Fasc_2511-25_2509-25.docx")

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x59, 0x59, 0x59)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HDR_FILL = "1F3A5F"
SUBHDR_FILL = "D9E1F2"
ZEBRA_FILL = "F2F5FA"

doc = Document()

# ---------- stili base ----------
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.08

for lvl, sz, col in [("Heading 1", 14, NAVY), ("Heading 2", 12, NAVY), ("Heading 3", 11, GREY)]:
    st = doc.styles[lvl]
    st.font.name = "Calibri"
    st.font.size = Pt(sz)
    st.font.color.rgb = col
    st.font.bold = True
    st.paragraph_format.space_before = Pt(10)
    st.paragraph_format.space_after = Pt(4)


def set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def cell_text(cell, text, bold=False, color=None, size=9.5, align=None, fill=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    runs = text.split("\n")
    for i, line in enumerate(runs):
        if i > 0:
            p.add_run().add_break()
        r = p.add_run(line)
        r.bold = bold
        r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
    if fill:
        set_cell_bg(cell, fill)


def add_para(text, bold=False, italic=False, size=10.5, color=None, align=None, space_after=4):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(space_after)
    return p


def add_bullet(text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
        r.font.size = Pt(10.5)
        r2 = p.add_run(text)
        r2.font.size = Pt(10.5)
    else:
        r = p.add_run(text)
        r.font.size = Pt(10.5)
    return p


def make_table(headers, rows, widths=None, header_fill=HDR_FILL, header_color=WHITE, zebra=True, font=9.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        cell_text(hdr[i], h, bold=True, color=header_color, size=font, align=WD_ALIGN_PARAGRAPH.CENTER, fill=header_fill)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        fill = ZEBRA_FILL if (zebra and ri % 2 == 0) else None
        for ci, val in enumerate(row):
            bold = (ci == 0 and len(headers) <= 2)
            cell_text(cells[ci], val, bold=bold, size=font, fill=fill)
    if widths:
        for row in t.rows:
            for ci, w in enumerate(widths):
                row.cells[ci].width = Cm(w)
    return t


# ====================================================================
# HEADER / FOOTER ricorrenti
# ====================================================================
section = doc.sections[0]
section.top_margin = Cm(2.6)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2.0)
section.right_margin = Cm(2.0)
section.header_distance = Cm(0.8)
section.footer_distance = Cm(0.8)

header = section.header
htbl = header.add_table(rows=1, cols=3, width=Cm(17))
htbl.style = "Table Grid"
hc = htbl.rows[0].cells
cell_text(hc[0], "Redatto a cura del:\nServizio Prevenzione e Protezione\nArsenale M.M. Taranto", size=8, align=WD_ALIGN_PARAGRAPH.LEFT)
cell_text(hc[1], "DOCUMENTO DI\nINFORMAZIONE ALLE DITTE\n- D.Lgs. 81/2008 Art. 26 comma 1 lett. b -", bold=True, size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER, color=NAVY)
cell_text(hc[2], "29 Maggio 2026\nBacino Ferrati (C22B)\nFasc. 2511/25 - 2509/25", size=8, align=WD_ALIGN_PARAGRAPH.RIGHT)

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = fp.add_run("SERVIZIO PREVENZIONE E PROTEZIONE - ARSENALE M.M. TARANTO")
fr.font.size = Pt(8)
fr.font.color.rgb = GREY

# ====================================================================
# FRONTESPIZIO
# ====================================================================
for _ in range(2):
    doc.add_paragraph()
add_para("ARSENALE MILITARE MARITTIMO", bold=True, size=20, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
add_para("TARANTO", bold=True, size=18, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
add_para("MARINARSEN", bold=False, size=11, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)

add_para("DOCUMENTO DI INFORMAZIONE ALLE DITTE", bold=True, size=16, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
add_para("ai sensi dell'art. 26, comma 1, lettera b) del D.Lgs. 81/2008 e s.m.i.", italic=True, size=11, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)

add_para("OGGETTO", bold=True, size=12, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
add_para(
    "ATTIVITA' DI MESSA IN SICUREZZA, TAGLIO E RIMOZIONE DELLE BARCHE PORTA GP65 E GP59 "
    "E ATTIVITA' DI PULIZIA DELLA PLATEA DA FANGHI E ACQUE PRESSO IL BACINO FERRATI (C22B) "
    "DELL'ARSENALE MILITARE MARITTIMO DI TARANTO",
    bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18,
)

# tabella riferimenti frontespizio
ft = doc.add_table(rows=0, cols=2)
ft.style = "Table Grid"
ft.alignment = WD_TABLE_ALIGNMENT.CENTER
front_rows = [
    ("Committente / Datore di Lavoro", "Arsenale Militare Marittimo di Taranto (MARINARSEN) - C.A. Alessandro BATTAGLIA"),
    ("Riferimenti contrattuali", "Fasc. 2511/25 Ordine nr. 1 (taglio e rimozione barche porta GP65 e GP59)\nFasc. 2509/25 Ordine nr. 6 (pulizia platea da fanghi e acque)\nA.Q. n. 327 di Rep. del 28/01/2025 di MARIUGCRA - Programma MCO Bacini di carenaggio"),
    ("Area di lavoro", "Bacino Ferrati (C22B) - platea, gradoni e banchine"),
    ("Coordinatore per la Sicurezza (CSP/CSE)", "Ing. Vincenzo MARASCIULO (Societa' Italiana S.r.l.)"),
    ("Inizio attivita' previsto", "03 Giugno 2026 (pulizia platea) - 08 Giugno 2026 (taglio barche porta)"),
    ("Redazione", "Servizio Prevenzione e Protezione - Arsenale M.M. Taranto"),
    ("Data / Revisione", "29 Maggio 2026 - Rev. 00"),
]
for k, v in front_rows:
    cells = ft.add_row().cells
    cell_text(cells[0], k, bold=True, size=10, fill=SUBHDR_FILL)
    cell_text(cells[1], v, size=10)
ft.columns[0].width = Cm(5.5)
ft.columns[1].width = Cm(11.5)

doc.add_paragraph()
add_para("Il Responsabile del Servizio Prevenzione e Protezione", size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
add_para("F.S.T. Ing. Giancarlo CAFORIO", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
add_para("_______________________________________", size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)

doc.add_page_break()

# ====================================================================
# 0 - PREMESSA
# ====================================================================
doc.add_heading("PREMESSA", level=1)
add_para(
    "Lo scopo del presente documento e' quello di fornire le informazioni, ai sensi dell'art. 26 comma 1 "
    "lettera b) del D.Lgs. 81/2008 e s.m.i., relative ai luoghi di lavoro oggetto del contratto e dipendenti "
    "dal Datore di lavoro di Marinarsen Taranto, nonche' gli elementi essenziali di conoscenza "
    "dell'organizzazione interna dell'Arsenale Militare Marittimo di Taranto."
)
add_para(
    "Le attivita' ricadono nel campo di applicazione del Titolo IV del D.Lgs. 81/2008 (cantieri temporanei o "
    "mobili), con Coordinamento della Sicurezza in fase di Progettazione ed Esecuzione (CSP/CSE) affidato "
    "all'Ing. Vincenzo Marasciulo della Societa' Italiana S.r.l. Le lavorazioni costituiscono la prosecuzione, "
    "per scelta condivisa, di quelle gia' effettuate per la pulizia del bacino dalle taccate in legno e in "
    "calcestruzzo. Il presente documento informativo integra e non sostituisce il Piano di Sicurezza e "
    "Coordinamento (PSC_Ferrati Rev. 5) e i Piani Operativi di Sicurezza (POS) delle singole imprese."
)
add_para(
    "Ai sensi dell'art. 26 comma 3 del D.Lgs. 81/2008, le Ditte affidatarie ed esecutrici sono tenute a "
    "partecipare alle riunioni di coordinamento indette dal CSE e a prendere visione del PSC e del relativo "
    "cronoprogramma, il cui aggiornamento e' obbligatorio al mutare delle condizioni operative o all'insorgere "
    "di nuovi rischi. Ogni variazione logistica, tecnica o di personale dovra' essere preventivamente "
    "comunicata al DEC, al CSE e al RSPP."
)

# ====================================================================
# 1 - RIFERIMENTI AL CONTRATTO
# ====================================================================
doc.add_heading("1) RIFERIMENTI AL CONTRATTO", level=1)
add_para(
    "Adesione all'Accordo Quadro (A.Q.) n. 327 di Repertorio del 28/01/2025 di MARIUGCRA - "
    "\"Programma di Mantenimento Condizioni Operative (MCO) per i Bacini di carenaggio di MARINARSEN Taranto\".",
)
add_bullet("taglio, riduzione volumetrica e rimozione delle barche porta GP65 e GP59 dal Bacino Ferrati, comprese le opere propedeutiche e di messa in sicurezza.", bold_lead="Fasc. 2511/25 - Ordine nr. 1: ")
add_bullet("attivita' di pulizia della platea del Bacino Ferrati da fanghi e acque, con impianto di trattamento in banchina (ciclo chiuso dell'acqua).", bold_lead="Fasc. 2509/25 - Ordine nr. 6: ")
add_para(
    "NOTA: gli allegati \"Lett. Fasc. 2511/25 Ord. 2\" e \"Allegato Fasc. 2511/25 Ord. 2\" presenti agli atti "
    "si riferiscono all'Ordine nr. 2 (revisione verricelli del Bacino Brin) e NON alle attivita' del Bacino "
    "Ferrati oggetto del presente documento: vanno pertanto sostituiti con la lettera d'ordine e l'allegato "
    "economico relativi al Fasc. 2511/25 Ordine nr. 1.",
    italic=True, size=9.5, color=GREY,
)

# ====================================================================
# 2 - DESCRIZIONE LAVORAZIONI E IMPRESE
# ====================================================================
doc.add_heading("2) DESCRIZIONE DELLE LAVORAZIONI E IMPRESE COINVOLTE", level=1)
add_para(
    "Nella platea del Bacino Ferrati sono presenti due barche porta che saranno oggetto di demolizione: la "
    "GP59, collassata e adagiata in posizione orizzontale sul fondo del bacino, e la GP65, appoggiata sui "
    "gradoni in posizione inclinata quasi verticale. La GP65 si trova nella posizione attuale da diversi mesi, "
    "con baricentro molto basso, in configurazione di equilibrio consolidata; non e' prevedibile alcun "
    "movimento relativo tra le due barche porta ne' scivolamento della GP65."
)

doc.add_heading("2.1 Imprese e ruoli", level=2)
make_table(
    ["Impresa", "Ruolo", "Attivita' principali"],
    [
        ["FAMAG S.r.l.", "Affidataria", "Tagli \"pronto forno\" propedeutici allo smaltimento; pulizia gradoni del Bacino Ferrati nelle zone occupate dalle barche porta GP65 e GP59; recupero blocchi di cemento residui in platea; lavaggio con idropulitrice; trasporto e smaltimento."],
        ["F&R S.r.l.\nVia Gesu' 21, 20121 Milano\nP.IVA 04688400276", "Subappaltatrice", "Progettazione esecutiva delle attivita' di taglio e rimozione delle strutture metalliche delle barche porta GP59 e GP65; valutazione e verifica step-by-step delle operazioni di riduzione volumetrica; direzione tecnica; emissione del progetto con fasi e procedure operative, messa in sicurezza e smaltimenti; taglio e rimozione delle barche porta, apprestamenti e mezzi (incluso ausilio di PLE)."],
        ["PEYRANI", "Noleggio a caldo", "Noleggio a caldo autogru (400 Ton e 200 Ton) per i sollevamenti delle sezioni tagliate."],
        ["OFFICINE JOLLY S.r.l.", "Esecutrice", "Prosecuzione della cantierizzazione gia' presente; realizzazione in banchina di un'area di deposito temporaneo per rottami e rifiuti (compresa installazione lamiere); prelievo dalla platea delle strutture metalliche e dei rifiuti prodotti e posizionamento nel deposito temporaneo; recupero legname residuo, lavaggio con idropulitrice, trasporto e smaltimento."],
        ["G-TEK", "Esecutrice", "Montaggio e mantenimento del ponteggio tipo torre scalo, idoneo anche al recupero dell'infortunato con barella; comodato d'uso alle imprese con limite operativo 150 kg/m2; revisione ogni 15 gg o dopo eventi meteo avversi."],
        ["STUDIO LO SASSO", "Esecutrice", "Rilievo topografico quotidiano con strumentazione dedicata, atto a rilevare eventuali movimenti o spostamenti indesiderati delle barche porta."],
        ["IDROVELOX di Petrelli Franco & Figli S.r.l.", "Esecutrice (Fasc. 2509/25)", "Pulizia della platea da fanghi e acque mediante impianto di trattamento in banchina e sistema di pompaggio (separazione fase solida/liquida, ciclo chiuso dell'acqua, smaltimento fanghi); impiego di piccoli mezzi movimento terra. Durata circa 20 gg."],
    ],
    widths=[4.0, 3.0, 10.0], font=9,
)

doc.add_heading("2.2 Fasi operative - Taglio e rimozione barche porta (Fasc. 2511/25)", level=2)
fasi = [
    ("FASE 1 - Messa in sicurezza spigoli di appoggio", "Incremento della sicurezza statica del punto di appoggio \"2\" della GP65 mediante posizionamento di blocchi in cemento armato (1500x1200x600 mm) impilati in due file da 5 blocchi, in modo da impedire qualsiasi movimento del punto di appoggio sul fondo del bacino."),
    ("FASE 2.1 - Posizionamento autogru e area di riduzione volumetrica", "Sezionamento iniziale della GP65 con taglio di elementi di peso compreso tra 5 e 20 t. Area di riduzione volumetrica protetta da lamiere in acciaio. Autogru 400 Ton in posizione \"1\" per le sezioni maggiori della GP65, poi autogru 200 Ton per le rimanenti; la 400 Ton si sposta in posizione \"2\" per la GP59."),
    ("FASE 2.2 - Progressione dei tagli (GP65)", "Decostruzione per sezioni successive con progressione controllata: direzione verticale dall'alto verso il basso, orizzontale dagli estremi verso il centro, mantenimento del nucleo centrale fino all'ultima fase. Ogni sezione e' preventivamente imbragata all'autogru, parzialmente tagliata, separata solo sotto tiro, sollevata e ridotta a terra. Tagli eseguiti da n. 2 operatori su PLE, posizionati su lati opposti, esterni alla traiettoria dei possibili movimenti delle sezioni."),
    ("FASE 3 - Riduzione a sezioni della GP59", "Riduzione delle sezioni a misura idonea ai sollevamenti con le due autogru. La GP59, gia' in posizione orizzontale sul fondo, non presenta problematiche di staticita'. La demolizione della GP59 sara' autorizzata solo quando l'assetto della GP65 sara' stabilizzato e sara' certa l'impossibilita' di ulteriori movimentazioni."),
    ("FASE 4 - Svuotamento completo del bacino", "Rimozione totale dei rottami ferrosi, pulizia e ripristino delle condizioni di sicurezza."),
]
for titolo, testo in fasi:
    add_bullet(testo, bold_lead=titolo + ": ")
add_para(
    "A vantaggio di sicurezza sono previsti: installazione di corpi morti e zavorre sull'appoggio libero come "
    "forza resistente allo scivolamento; rilievo topografico quotidiano di punti scelti per verificare "
    "eventuali spostamenti, con eventuale modifica del piano di azione. Tutti i tagli comportano un "
    "abbassamento del baricentro e una riduzione del peso che spinge verso lo scivolamento.",
    italic=True, size=9.5, color=GREY,
)

doc.add_heading("2.3 Pulizia platea da fanghi e acque (Fasc. 2509/25)", level=2)
add_para(
    "L'attivita' viene effettuata dall'impresa Idrovelox mediante un impianto di trattamento installato in "
    "banchina. Il fango viene prelevato con sistema di pompaggio e trattato in modo che l'acqua risulti "
    "riutilizzabile per aspirare altro fango (ciclo chiuso dell'acqua), riducendo la produzione di materiale da "
    "smaltire. I fanghi sono accumulati in cassone scarrabile in banchina e successivamente smaltiti. E' "
    "prevista la presenza in cantiere di piccoli mezzi movimento terra. Durata stimata circa 20 giorni."
)

# ====================================================================
# 3 - FIGURE RESPONSABILI
# ====================================================================
doc.add_heading("3) FIGURE RESPONSABILI AI SENSI DEL D.LGS. 81/2008", level=1)
add_para(
    "Marinarsen ha una struttura gerarchica definita con Decreto Ministeriale (DM) e disposizioni interne. "
    "Di seguito i nominativi del Datore di lavoro, dei dirigenti e dei preposti interessati alla presente attivita':"
)
make_table(
    ["Ruolo", "Nominativo", "Recapito"],
    [
        ["Datore di Lavoro - Direttore p.t. dell'Arsenale di Taranto", "C.A. Alessandro BATTAGLIA", "-"],
        ["Capo Reparto Manutenzioni Navali / RUP", "C.V. (AN) Marco ACCOTO", "099 775 3940"],
        ["Capo Sezione Bacini (Dirigente alla Sicurezza)", "C.F. Marco BONATTO", "099 775 2335"],
        ["DEC Fascicolo", "C.F. Nicola BOCCARDI", "099 775 7557"],
        ["Capo Sezione Reti Elettriche", "F.T. Angelo CARDELLICCHIO", "099 775 3020"],
        ["Responsabile Servizio Prevenzione e Protezione", "F.S.T. Ing. Giancarlo CAFORIO", "099 775 3103"],
        ["Medico Competente", "Dott. Giovanni TRIA", "099 775 4028"],
        ["Coordinatore Sicurezza Progettazione/Esecuzione (CSP/CSE)", "Ing. Vincenzo MARASCIULO - Societa' Italiana S.r.l.", "099 400 5718"],
    ],
    widths=[7.0, 6.5, 3.5], font=9.5,
)

doc.add_heading("Rappresentanti dei Lavoratori per la Sicurezza (RLS)", level=2)
add_para("Personale Civile: Alessandro SCIALPI - Giovanni CHIFFI - Ignazio BARBUTO - Pietro AVELLINO - Saverio CAPODIFERRO - Valentina FALCONE.", size=10)

# ====================================================================
# 4 - INFORMAZIONI GENERALI
# ====================================================================
doc.add_heading("4) INFORMAZIONI GENERALI SULL'ARSENALE DI TARANTO", level=1)
add_bullet("ARSENALE MILITARE MARITTIMO TARANTO (detto anche MARINARSEN TARANTO).", bold_lead="Denominazione Ente: ")
add_bullet("Taranto, Piazza Ammiraglio Pasquale Leonardi - CAP 74123.", bold_lead="Ubicazione: ")
add_bullet("MARIVIGILANZA - AREA SUD - TARANTO.", bold_lead="Organo di Vigilanza Competente: ")

# ====================================================================
# 5 - EMERGENZE
# ====================================================================
doc.add_heading("5) GESTIONE EMERGENZE, EVACUAZIONE E PRIMO SOCCORSO", level=1)
add_para(
    "In caso di emergenza/evacuazione raggiungere il centro di raccolta esterno all'area di lavoro indicato "
    "nella planimetria di cantiere (Layout di cantiere allegato). Mantenere sempre libere le vie di accesso "
    "per consentire l'agevole transito dei mezzi di soccorso."
)
doc.add_heading("Chiamate di emergenza", level=2)
make_table(
    ["Tipologia di emergenza", "Recapito"],
    [
        ["EMERGENZA INCENDIO", "115 - Vigili del Fuoco"],
        ["EMERGENZA SANITARIA", "118 - Pronto Soccorso"],
        ["SALA MEDICA ARSENALE", "099 775 2841 (lun-gio 08:00-16:00; ven 08:00-12:00)"],
        ["EMERGENZA ELETTRICA - Guardia Circuiti Sezione Reti Elettriche", "tel. militare 22604 - cellulare 099 775 2604"],
        ["EMERGENZA BACINI FERRATI (orario 08:00-13:00)", "tel. militare 23765 / 23762 - cell. 099 775 3765 / 099 775 3762"],
        ["EMERGENZA BACINI FERRATI (turno guardia 13:00-08:00)", "tel. militare 22916 - cellulare 099 775 2916"],
    ],
    widths=[9.0, 8.0], font=9.5, zebra=True,
)
add_para(
    "Recupero infortunato: deve essere sempre garantita la disponibilita' di una gru per il tempestivo "
    "recupero in sicurezza di eventuali lavoratori infortunati in platea. Il ponteggio tipo torre scalo "
    "(G-Tek) e' idoneo al recupero dell'infortunato con barella.",
    bold=False, size=10,
)

doc.add_heading("6) INFORMAZIONE SUGLI ACCESSI NEL COMPRENSORIO", level=2)
add_para(
    "L'accesso al comprensorio e' regolato dalle vigenti Comunicazioni di Servizio. Tutto il personale deve "
    "essere identificabile mediante tesserino di riconoscimento con fotografia, formato e aggiornato alla "
    "sicurezza con i corsi previsti dalla legge. Il transito veicolare dei mezzi pesanti deve essere effettuato "
    "da Porta Levante."
)

# ====================================================================
# 7 - RISCHI
# ====================================================================
doc.add_page_break()
doc.add_heading("7) INFORMAZIONI SUI RISCHI", level=1)
doc.add_heading("7.1 Codificazione dei fattori di rischio", level=2)
make_table(
    ["Codice fattore di rischio", "Ambiti / luoghi di potenziale rischio"],
    [
        ["1 - Agenti Biologici", "Materiali contaminati, rottami metallici, sedimenti e acque stagnanti in platea: rischio di infezione da ferite/abrasioni."],
        ["2 - Agenti Chimici", "Residui di idrocarburi, oli, solventi e sostanze tossiche/infiammabili nei sedimenti e nelle acque; vapori e aerosol."],
        ["3 - Caduta dall'alto / Lavori in quota", "Operazioni di taglio su PLE, accesso ai gradoni e al ponteggio torre scalo."],
        ["10 - Apparecchi di sollevamento", "Autogru 400 Ton e 200 Ton, PLE; movimentazione sezioni metalliche 5-20 t."],
        ["11 - Movimentazione meccanica / mezzi", "Camion con ragno, piccoli mezzi movimento terra, carico/scarico rottami."],
        ["13 - Impianti Elettrici", "Elettrocuzione; alimentazione attrezzature di taglio e pompe di sentina."],
        ["14 - Primo Soccorso", "Recupero infortunato in platea e da PLE/ponteggio."],
        ["15 - Rischio Incendio", "Tagli a caldo (ossitaglio/cesoie) su strutture metalliche con possibili residui combustibili."],
        ["17 - Interconnessione e viabilita'", "Strade, banchine, marciapiedi, scale, sottoservizi del comprensorio; transito mezzi."],
        ["18 - Locali di Lavoro", "Bacino Ferrati (platea, gradoni e banchine): caduta dai gradoni, scivolamento, allagamento."],
        ["19 - Rischio strutturale / scivolamento barche porta", "Possibile movimento/scivolamento delle barche porta durante le fasi di taglio."],
        ["23 - Organizzazione del Lavoro", "Interferenze tra imprese, coordinamento, gestione rifiuti, formazione del personale."],
    ],
    widths=[5.0, 12.0], font=9.5,
)

doc.add_heading("7.2 Quadro di sintesi dei rischi e misure", level=2)

risk_blocks = [
    ("1 - Rischio Biologico",
     "La presenza di materiali contaminati, rottami metallici e sedimenti inquinati aumenta la probabilita' di "
     "ferite da taglio o abrasioni e di conseguenti contaminazioni; l'elevata umidita' e la scarsa aerazione "
     "favoriscono la proliferazione di microrganismi. Dotare il personale di guanti, sovratute e calzature "
     "impermeabili; vietare cibo/bevande/fumo nelle aree di intervento; disinfettare immediatamente qualsiasi "
     "ferita; sorveglianza sanitaria con vaccinazione antitetanica in corso di validita'."),
    ("2 - Rischio Chimico",
     "Rischio legato a residui di idrocarburi, oli, solventi e sostanze tossiche/infiammabili nei sedimenti e "
     "nelle acque stagnanti. Prevedere monitoraggio delle atmosfere di lavoro, DPI specifici (maschere filtranti, "
     "guanti e tute impermeabili), procedure di decontaminazione e trasmissione al DEC delle schede di sicurezza "
     "(SDS) dei prodotti utilizzati."),
    ("3/18 - Caduta dall'alto e lavori in quota",
     "Per i tagli in quota gli operatori operano esclusivamente da n. 2 PLE marcate CE, condotte da personale "
     "abilitato e con idoneita' sanitaria. Sulla PLE salgono solo il manovratore abilitato e gli operatori "
     "incaricati, dotati di DPI anticaduta (casco, imbracatura EN 361 con dispositivo anticaduta). Per lavori in "
     "quota superiori a 2 m si applicano integralmente le prescrizioni del Titolo IV, Capo II del D.Lgs. 81/2008. "
     "Vietato l'uso di trabattelli per quote superiori a 8 m nei luoghi esterni."),
    ("10 - Apparecchi di sollevamento - movimentazione carichi",
     "Prima dell'utilizzo l'impresa noleggiatrice (Peyrani) deve consegnare: dichiarazione di conformita' CE, "
     "registro di uso/manutenzione, verbale dell'ultima verifica periodica, dichiarazione di regolare "
     "manutenzione, certificazione CE degli accessori di sollevamento e attestato di abilitazione del "
     "conduttore. Vietato posizionare gli stabilizzatori senza piastre o su piani non resistenti/grigliati/"
     "cunicoli. Non sono note le portate massime di carico in banchina. Vietato movimentare carichi con "
     "personale nell'area sottostante: interdire la zona fino a 10 m da ambo i lati del tragitto del pescante. "
     "Comunicazione continua gruista-operatore a terra (gestuale codificata o radio); responsabile unico della "
     "movimentazione."),
    ("13 - Impianti Elettrici",
     "La posa dei cavi non deve interferire con il camminamento dei lavoratori e la movimentazione dei carichi. "
     "Tutte le apparecchiature devono essere conformi CE e norme CEI, con prolunghe industriali integre e "
     "protezione differenziale ad alta sensibilita' (30 mA), idonee all'ambiente umido. Misure ed eventuali "
     "interventi in presenza di rischio elettrico solo da personale abilitato PES/PAV (CEI 11-27)."),
    ("14 - Primo Soccorso",
     "Deve essere sempre presente una gru per il tempestivo recupero in sicurezza di eventuali infortunati. Il "
     "personale incaricato sara' formato e informato sulle procedure di intervento d'emergenza. Il ponteggio "
     "torre scalo e' utilizzabile per il recupero dell'infortunato con barella."),
    ("15 - Rischio Incendio (tagli a caldo)",
     "Prima dell'inizio attivita' verificare l'assenza di vapori infiammabili residui. L'impianto antincendio "
     "del Bacino Ferrati non e' disponibile: la Ditta esecutrice deve dotarsi di estintori supplementari idonei "
     "posizionati in prossimita' dell'area di lavoro. Durante le operazioni di taglio che generano scintille o "
     "proiezione di materiale incandescente, predisporre schermi/protezioni ed estintori di pronto impiego. "
     "Vietato l'uso di fiamme libere non autorizzate."),
    ("17 - Interconnessione e viabilita'",
     "Non sono noti i carichi massimi sopportabili delle massicciate ne' la distribuzione dei sottoservizi. "
     "Limite di velocita' 20 km/h sulle strade interne. Altezza massima veicolare 4,25 m (4,50 m nel tratto "
     "Porta Levante - Piazzale Rottami). Attenzione a buche, beole danneggiate, tombini e grigliati sconnessi. "
     "Il transito pedonale e' consentito solo nelle aree e sui marciapiedi previsti; le banchine sul mare non "
     "sono delimitate (pericolo di caduta in mare). Non ingombrare vie di evacuazione e accessi ai presidi "
     "antincendio."),
    ("18/19 - Locali di lavoro e rischio scivolamento barche porta",
     "Il Bacino Ferrati presenta pericolo di caduta dai gradoni, scivolamento, caduta di materiale dall'alto, "
     "inciampo, interferenza con altri cantieri e allagamento in caso di eventi catastrofici. E' vietato "
     "l'ingresso ai non addetti. E' previsto il monitoraggio topografico quotidiano (Studio Lo Sasso) per "
     "rilevare spostamenti indesiderati delle barche porta; in caso di movimenti non attesi le attivita' "
     "vengono sospese e il piano di azione rivisto dal CSE. Nessun operatore puo' accedere alla platea o ad "
     "aree non autorizzate senza autorizzazione scritta e senza accompagnamento di un rappresentante di "
     "Marinarsen (DEC, C.F. Boccardi)."),
    ("23 - Organizzazione del lavoro",
     "Tutte le attivita' sono supervisionate da un referente del Bacino in stretto coordinamento con il DEC e "
     "il CSE. Il personale deve essere formato, identificabile con tesserino e informato sulle prescrizioni del "
     "presente documento. Lo stoccaggio di materiali e rifiuti deve essere ordinato, segnalato e non di "
     "intralcio; pulizia e rimozione quotidiana dei residui. Transennare e segnalare sempre le aree di lavoro. "
     "Vietato l'uso di bevande alcoliche, in particolare per gruisti, manovratori di PLE e addetti ai lavori in "
     "quota. I rifiuti vanno smaltiti presso centri autorizzati (D.Lgs. 152/2006) con consegna al DEC della 4^ "
     "copia del FIR indicante il fascicolo e i quantitativi."),
]
for titolo, testo in risk_blocks:
    add_para(titolo, bold=True, size=11, color=NAVY, space_after=1)
    add_para(testo, size=10)

# ====================================================================
# 8 - RISCHI INTERFERENZIALI
# ====================================================================
doc.add_heading("8) RISCHI INTERFERENZIALI E MISURE DI COORDINAMENTO", level=1)
add_para(
    "Considerata la compresenza di piu' imprese (Famag, F&R, Officine Jolly, Peyrani, G-Tek, Studio Lo Sasso, "
    "Idrovelox) nella stessa area, il CSE effettua l'analisi delle interferenze e predispone/aggiorna il "
    "cronoprogramma. Lo scrivente ufficio del CSE e' presente in cantiere ogni giorno per almeno 3 ore e redige "
    "quotidianamente il verbale con le prescrizioni di sicurezza, che ha valore di aggiornamento del PSC."
)
make_table(
    ["Descrizione del rischio interferenziale", "Misure di prevenzione e provvedimenti da adottare"],
    [
        ["Scivolamenti, aperture e/o ostacoli non segnalati sui camminamenti", "Segnalare con specifica segnaletica le superfici di transito a rischio scivolamento; individuare e segnalare ostacoli e aperture."],
        ["Esecuzione lavori con presenza di altre ditte e/o personale", "Informare preventivamente i responsabili delle imprese interferenti; fornire informazioni a tutto il personale; attenersi alle indicazioni del CSE e alle sequenze del cronoprogramma."],
        ["Carichi sospesi, carichi mobili e caduta di oggetti dall'alto", "Limitare la sospensione dei carichi ai tempi strettamente necessari; vietare il passaggio sotto i carichi sospesi; interdire l'area fino a 10 m da ambo i lati del tragitto del pescante; usare DPI per la protezione del capo."],
        ["Produzione di schegge, polveri, fumi ed esalazioni in presenza di personale estraneo", "Informare preventivamente i responsabili segnalando il pericolo; interdire il transito agli estranei nella zona interessata; ove possibile spostare le lavorazioni in orari extralavorativi."],
        ["Presenza di impianti elettrici sotto tensione", "Non lasciare cavi volanti nelle zone di passaggio; non eseguire manutenzioni di propria iniziativa; non sovraccaricare l'impianto; impiegare dispositivi dielettrici; segnalare quadri e impianti in manutenzione."],
        ["Presenza di macchinari da taglio / pressatura meccanica", "Informare preventivamente i responsabili; interdire il transito agli estranei nella zona interessata; predisporre protezioni apposite."],
        ["Impiego di mezzi mobili e veicoli (camion, autogru, PLE)", "Segnalare la presenza di mezzi in movimento; vietare l'avvicinamento del personale non interessato; transennare la zona di lavoro; rispettare i limiti di velocita'."],
        ["Utilizzo del ponteggio (torre scalo)", "Consentito il solo utilizzo; segnalare al preposto anomalie (giunti non serrati, elementi danneggiati) e non utilizzarlo fino al ripristino; non modificare/manomettere; tenere chiuse le botole salvo transito; rispettare il carico max 150 kg/m2."],
        ["Presenza in cantiere di personale non dipendente delle ditte esecutrici", "Accesso previa autorizzazione del preposto/responsabile di cantiere e verifica delle lavorazioni in corso; obbligo di elmetto e calzature antinfortunistiche per tutti gli autorizzati."],
        ["Rischio scivolamento/instabilita' barche porta", "Sfasamento temporale delle lavorazioni: la demolizione della GP59 e' autorizzata solo a GP65 stabilizzata; monitoraggio topografico quotidiano; sospensione immediata in caso di spostamenti non attesi."],
    ],
    widths=[6.5, 10.5], font=9,
)
add_para(
    "Sfasamento spaziale e temporale: le attivita' di pulizia platea (Idrovelox) e le attivita' di taglio "
    "(Famag/F&R) nelle stesse zone non possono essere eseguite in contemporanea sovrapposizione; le sequenze "
    "operative sono definite quotidianamente dal CSE nel verbale di coordinamento.",
    italic=True, size=9.5, color=GREY,
)

# ====================================================================
# 9 - DPI
# ====================================================================
doc.add_heading("9) DISPOSITIVI DI PROTEZIONE INDIVIDUALE (DPI) OBBLIGATORI", level=1)
make_table(
    ["Rischio", "DPI obbligatori"],
    [
        ["Generale accesso cantiere", "Elmetto di protezione (EN 397); calzature di sicurezza con puntale in acciaio e a sfilamento rapido (EN ISO 20345); indumenti ad alta visibilita' (EN ISO 20471)."],
        ["1 - Agenti Biologici", "Guanti impermeabili; sovratuta/tuta impermeabile; stivali impermeabili con puntale; maschera semifacciale con filtro P3 in presenza di aerosol."],
        ["2 - Agenti Chimici", "Guanti resistenti a oli/idrocarburi (EN 374); maschera filtrante (filtro combinato A2P3 in presenza di vapori organici); tuta impermeabile."],
        ["3 - Lavori in quota / PLE", "Imbracatura anticaduta (EN 361) con dispositivo anticaduta; casco con sottogola."],
        ["10/11 - Sollevamento e mezzi", "Elmetto; guanti per rischio meccanico; calzature di sicurezza; alta visibilita' per il personale a terra."],
        ["13 - Impianti elettrici", "Guanti isolanti (per lavori autorizzati); calzature isolanti (EN 50321); occhiali di sicurezza."],
        ["15 - Taglio a caldo", "Visiera/occhiali per saldatura; guanti e grembiule per saldatore; calzature di sicurezza; estintore di pronto impiego nelle vicinanze."],
        ["17 - Lavoro vicino all'acqua", "Giubbotto salvagente per chi opera in prossimita' dell'acqua/banchine non delimitate; calzature con suola antiscivolo."],
    ],
    widths=[4.5, 12.5], font=9.5,
)

# ====================================================================
# 10 - CRONOPROGRAMMA
# ====================================================================
doc.add_heading("10) CRONOPROGRAMMA SINTETICO", level=1)
make_table(
    ["Periodo", "Attivita'", "Impresa"],
    [
        ["dal 03/06/2026 (circa 20 gg)", "Pulizia platea da fanghi e acque - impianto di trattamento in banchina", "Idrovelox (Fasc. 2509/25)"],
        ["dal 03/06/2026", "Prosecuzione cantierizzazione, area deposito temporaneo, mantenimento ponteggio torre scalo", "Officine Jolly / G-Tek"],
        ["dall'08/06/2026", "Messa in sicurezza appoggi e avvio taglio/riduzione GP65", "Famag / F&R / Peyrani"],
        ["a seguire", "Riduzione e rimozione GP59 (previa stabilizzazione GP65)", "Famag / F&R / Peyrani"],
        ["quotidiano", "Rilievo topografico di controllo movimenti barche porta", "Studio Lo Sasso"],
        ["finale", "Svuotamento completo del bacino, rimozione rottami, pulizia e ripristino", "Officine Jolly / Famag"],
    ],
    widths=[4.5, 9.0, 3.5], font=9.5,
)

# ====================================================================
# 11 - DISPOSIZIONI FINALI
# ====================================================================
doc.add_heading("11) DISPOSIZIONI FINALI E PRESCRIZIONI GENERALI", level=1)
for t in [
    "E' vietato l'ingresso all'area di cantiere ai non addetti ai lavori; transennare e segnalare sempre le aree di lavoro.",
    "Nessun operatore puo' accedere alla platea o ad aree non autorizzate senza autorizzazione scritta e senza accompagnamento di un rappresentante di Marinarsen (DEC).",
    "Tutto il personale deve essere formato e aggiornato alla sicurezza, identificabile con tesserino con fotografia e informato sulle prescrizioni del presente documento.",
    "I numeri di emergenza e le figure responsabili devono essere riportati nella segnaletica di cantiere.",
    "Il preposto di ciascuna ditta garantisce l'osservanza degli obblighi di legge, l'uso dei DPI e segnala tempestivamente al DEC, al CSE e al RSPP qualsiasi anomalia, near-miss o infortunio.",
    "E' obbligatorio mantenere sempre libere le vie di accesso per il transito dei mezzi di soccorso.",
    "I rifiuti vanno smaltiti presso centri autorizzati (D.Lgs. 152/2006), con consegna al DEC della 4^ copia del FIR indicante fascicolo e quantitativi.",
    "Le Ditte devono partecipare alle riunioni di coordinamento indette dal CSE e attenersi al cronoprogramma e ai verbali quotidiani di coordinamento.",
]:
    add_bullet(t)

doc.add_paragraph()
add_para("ALLEGATI RICHIAMATI", bold=True, size=11, color=NAVY)
for a in [
    "Piano di Sicurezza e Coordinamento - PSC_Ferrati Rev. 5 (CSE Ing. V. Marasciulo)",
    "Layout di cantiere (recinzione, accessi mezzi/pedonale, parcheggio, area riduzione volumetrica)",
    "Layout 1 - Vista assonometrica ponteggio torre scalo (rev. 02)",
    "Layout 2 - Pianta, prospetto frontale e laterale ponteggio torre scalo (rev. 02)",
    "Notifica preliminare (da aggiornare con impresa F&R e oggetto dei lavori)",
    "Nomina CSE - Ing. Vincenzo Marasciulo",
    "Documentazione ponteggio tipo torre scalo (PiMUS n. 135_2025)",
]:
    add_bullet(a)

doc.add_paragraph()
tbl = doc.add_table(rows=1, cols=2)
tbl.style = "Table Grid"
c = tbl.rows[0].cells
cell_text(c[0], "Il Responsabile del Servizio Prevenzione e Protezione\n\nF.S.T. Ing. Giancarlo CAFORIO\n\n____________________________", size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
cell_text(c[1], "Il Coordinatore per la Sicurezza (CSP/CSE)\n\nIng. Vincenzo MARASCIULO\n\n____________________________", size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

doc.save(OUT)
print("[OK] documento generato:", OUT)
print("Pagine/elementi:", len(doc.paragraphs), "paragrafi,", len(doc.tables), "tabelle")
