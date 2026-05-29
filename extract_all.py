"""
Estrae il testo da tutti i documenti di riferimento per costruire il nuovo
Documento Informativo alle Ditte (Fascicoli 2511_25 e 2509_25 - Bacino Ferrati).
"""
import os
import sys
import pdfplumber
from docx import Document

BASE = "/projects/sandbox/ITALIC"
OUT = os.path.join(BASE, "_extracted")
os.makedirs(OUT, exist_ok=True)


def extract_pdf(path, out_name, max_pages=None):
    print(f"\n=== {os.path.basename(path)} ===")
    out_path = os.path.join(OUT, out_name)
    with pdfplumber.open(path) as pdf, open(out_path, "w", encoding="utf-8") as f:
        n = len(pdf.pages)
        f.write(f"# {os.path.basename(path)}\n")
        f.write(f"# {n} pagine\n\n")
        pages = pdf.pages if max_pages is None else pdf.pages[:max_pages]
        for i, page in enumerate(pages, 1):
            try:
                t = page.extract_text() or ""
            except Exception as e:
                t = f"[ERR pagina {i}: {e}]"
            f.write(f"\n----- PAGINA {i}/{n} -----\n")
            f.write(t)
            f.write("\n")
    print(f"  pages={n}  ->  {out_path}")
    return out_path


def extract_docx(path, out_name):
    print(f"\n=== {os.path.basename(path)} ===")
    out_path = os.path.join(OUT, out_name)
    d = Document(path)
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(f"# {os.path.basename(path)}\n\n")
        f.write("## STILI / TITOLI / PARAGRAFI\n")
        for p in d.paragraphs:
            style = p.style.name if p.style else ""
            txt = p.text
            if txt.strip():
                f.write(f"[{style}] {txt}\n")
        f.write("\n## TABELLE\n")
        for ti, table in enumerate(d.tables, 1):
            f.write(f"\n--- TAB {ti} ({len(table.rows)} righe x {len(table.columns)} col) ---\n")
            for row in table.rows:
                cells = [c.text.strip().replace("\n", " | ") for c in row.cells]
                f.write(" || ".join(cells) + "\n")
    print(f"  ->  {out_path}")
    return out_path


targets_pdf = [
    ("Fascicolo 2508.25 bacino FERRATI_Rev1_Firmato.pdf", "01_fasc2508.txt"),
    ("PSC_Ferrati_27042026 rev5.pdf", "02_psc_ferrati.txt"),
    ("Layout 1 - rev_02.pdf", "03_layout1.txt"),
    ("Layout 2 - rev_02.pdf", "04_layout2.txt"),
    ("Layout di cantiere.pdf", "05_layout_cant.txt"),
    ("Allegato Fasc. 2511_25 Ord 2.pdf", "06_allegato_2511.txt"),
    ("Lett. Fasc. 2511_25 Ord 2.pdf", "07_lett_2511.txt"),
]

for fname, out in targets_pdf:
    p = os.path.join(BASE, fname)
    if os.path.exists(p):
        try:
            extract_pdf(p, out)
        except Exception as e:
            print(f"  ERR: {e}")
    else:
        print(f"  MISSING: {p}")

docx_path = os.path.join(BASE, "20260521- Doc.info ditte - Fascicolo_6301_26_Mancarella_Rev_finale.docx")
if os.path.exists(docx_path):
    extract_docx(docx_path, "00_template_mancarella.txt")

print("\n[OK] estrazione completata in", OUT)
